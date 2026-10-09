"""Bounded XLSX ingestion; formulas and external links are never evaluated."""

from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import iterparse

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from backend.app.domain.errors import ApplicationError

MAX_CELLS = 1_000_000
MAX_EXPANDED_BYTES = 200 * 1024 * 1024


def validate_xlsx(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                raise ValueError("No contiene un libro Excel")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise ValueError("No se admiten macros")
            if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ApplicationError("file_too_large", "El Excel descomprimido supera el límite", status_code=413)
            if any(entry.flag_bits & 1 for entry in entries):
                raise ValueError("No se admiten libros cifrados")
            if archive.testzip() is not None:
                raise ValueError("El libro está dañado")
            cell_count = 0
            merged_cells = 0
            for name in names:
                if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
                    continue
                with archive.open(name) as xml:
                    for _event, element in iterparse(xml, events=("end",)):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if tag == "c":
                            cell_count += 1
                        elif tag == "mergeCell":
                            left, top, right, bottom = range_boundaries(element.attrib["ref"])
                            merged_cells += (right - left + 1) * (bottom - top + 1)
                        if cell_count + merged_cells > MAX_CELLS:
                            raise ApplicationError("file_too_large", "El Excel supera el límite de celdas", status_code=413)
                        element.clear()
        # Validate the package, not just its ZIP signature, before accepting it.
        with path.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            try:
                if sum((sheet.max_row or 0) * (sheet.max_column or 0) for sheet in workbook) > MAX_CELLS:
                    raise ApplicationError("file_too_large", "El Excel supera el límite de celdas", status_code=413)
            finally:
                workbook.close()
    except ApplicationError:
        raise
    except Exception as exc:
        raise ApplicationError("invalid_xlsx", "El archivo no contiene un libro XLSX válido") from exc


def read_workbook(path: Path):
    validate_xlsx(path)
    return load_workbook(path, data_only=False, keep_links=False)
