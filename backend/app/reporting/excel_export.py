"""Auditable Excel exports generated entirely by the backend."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from backend.app.config import Settings
from backend.app.infrastructure.database.models import (
    ApprovalEvent,
    CandidateFactor,
    ChemicalComposition,
    ClassificationCandidate,
    ClassificationResult,
    ClassificationRun,
    ClassificationSelection,
    Correction,
    EvidenceLink,
    Heat,
    Manufacturer,
    MillCertificate,
    Observation,
    Product,
    RuleSet,
)


class ExcelExportService:
    sheet_names = [
        "Resumen", "Actas", "Coladas", "Rollos", "Composición",
        "Clasificación", "Evidencia", "Auditoría",
    ]

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    @staticmethod
    def resolve_runs(
        session: Session, certificate_ids: list[int],
        classification_run_ids: list[int] | None, official: bool,
    ) -> list[ClassificationRun]:
        query = select(ClassificationRun).where(ClassificationRun.certificate_id.in_(certificate_ids))
        if classification_run_ids is not None:
            requested = set(classification_run_ids)
            runs = list(session.scalars(query.where(ClassificationRun.id.in_(requested)).order_by(ClassificationRun.id)))
            if len(runs) != len(requested):
                raise ValueError("Una ejecución no pertenece al alcance seleccionado")
        else:
            latest: dict[int, ClassificationRun] = {}
            for run in session.scalars(query.order_by(ClassificationRun.id.desc())):
                latest.setdefault(run.certificate_id, run)
            runs = list(latest.values())
        if official:
            if {run.certificate_id for run in runs} != set(certificate_ids):
                raise ValueError("Un reporte oficial requiere una ejecución aprobada por cada acta")
            if any(run.approval_status != "approved" for run in runs):
                raise ValueError("Un reporte oficial sólo admite ejecuciones aprobadas")
        return runs

    def export_certificates(
        self,
        session: Session,
        *,
        certificate_ids: list[int],
        heat_ids: list[int] | None = None,
        classification_run_ids: list[int] | None = None,
        destination: Path,
        official: bool,
    ) -> Path:
        if not certificate_ids:
            raise ValueError("Debe seleccionar al menos un acta")
        certificate_ids = list(dict.fromkeys(certificate_ids))
        requested_heat_ids = list(dict.fromkeys(heat_ids or []))

        certificates = session.scalars(
            select(MillCertificate).where(MillCertificate.id.in_(certificate_ids)).order_by(MillCertificate.id)
        ).all()
        if len(certificates) != len(set(certificate_ids)):
            raise ValueError("Una o más actas no existen")

        if official:
            unapproved_certs = [c.id for c in certificates if c.approval_status != "approved"]
            if unapproved_certs:
                raise ValueError(
                    f"Un reporte oficial sólo admite actas aprobadas (actas pendientes: {unapproved_certs})"
                )

        restrict_heats = bool(requested_heat_ids)
        heats_query = select(Heat).where(Heat.certificate_id.in_(certificate_ids))
        if restrict_heats:
            heats_query = heats_query.where(Heat.id.in_(requested_heat_ids))
        heats = session.scalars(heats_query.order_by(Heat.id)).all()
        if restrict_heats and len(heats) != len(requested_heat_ids):
            raise ValueError("Una o más coladas no pertenecen al alcance seleccionado")
        selected_heat_ids = [item.id for item in heats]

        products_query = select(Product).where(Product.certificate_id.in_(certificate_ids))
        if restrict_heats:
            products_query = products_query.where(Product.heat_id.in_(selected_heat_ids))
        products = session.scalars(products_query.order_by(Product.id)).all()
        product_ids = [item.id for item in products]

        observation_query = select(Observation).where(Observation.certificate_id.in_(certificate_ids))
        if restrict_heats:
            observation_query = observation_query.where(
                or_(Observation.heat_id.in_(selected_heat_ids), Observation.product_id.in_(product_ids),
                    (Observation.heat_id.is_(None) & Observation.product_id.is_(None)))
            )
        observations = session.scalars(observation_query.order_by(Observation.id)).all()
        obs_by_product: dict[int, int] = {}  # product_id -> row index in Evidencia sheet

        chemistry = []
        if selected_heat_ids or product_ids:
            chemistry = session.scalars(
                select(ChemicalComposition)
                .where(or_(ChemicalComposition.heat_id.in_(selected_heat_ids), ChemicalComposition.product_id.in_(product_ids)))
                .order_by(ChemicalComposition.id)
            ).all()

        runs = self.resolve_runs(session, certificate_ids, classification_run_ids, official)

        run_ids = [item.id for item in runs]
        results = session.scalars(
            select(ClassificationResult).where(
                ClassificationResult.classification_run_id.in_(run_ids),
                ClassificationResult.product_id.in_(product_ids),
            ).order_by(ClassificationResult.id)
        ).all() if run_ids else []
        result_ids = [item.id for item in results]

        candidates = session.scalars(
            select(ClassificationCandidate)
            .where(ClassificationCandidate.classification_result_id.in_(result_ids))
            .order_by(ClassificationCandidate.classification_result_id, ClassificationCandidate.rank)
        ).all() if result_ids else []
        candidates_by_result: dict[int, list[ClassificationCandidate]] = {}
        for cand in candidates:
            candidates_by_result.setdefault(cand.classification_result_id, []).append(cand)

        candidate_ids = [cand.id for cand in candidates]
        factors = session.scalars(
            select(CandidateFactor)
            .where(CandidateFactor.candidate_id.in_(candidate_ids))
            .order_by(CandidateFactor.candidate_id, CandidateFactor.sequence)
        ).all() if candidate_ids else []
        factors_by_candidate: dict[int, list[CandidateFactor]] = {}
        for factor in factors:
            factors_by_candidate.setdefault(factor.candidate_id, []).append(factor)

        factor_ids = [factor.id for factor in factors]
        factor_links = session.scalars(
            select(EvidenceLink)
            .where(EvidenceLink.candidate_factor_id.in_(factor_ids))
            .order_by(EvidenceLink.candidate_factor_id, EvidenceLink.id)
        ).all() if factor_ids else []
        rule_by_observation: dict[int, list[str]] = {}
        factor_by_id = {f.id: f for f in factors}
        for link in factor_links:
            if link.observation_id and link.candidate_factor_id in factor_by_id:
                codes = rule_by_observation.setdefault(link.observation_id, [])
                code = factor_by_id[link.candidate_factor_id].rule_code
                if code not in codes:
                    codes.append(code)

        selections = session.scalars(
            select(ClassificationSelection)
            .where(ClassificationSelection.classification_result_id.in_(result_ids))
            .order_by(ClassificationSelection.classification_result_id, ClassificationSelection.created_at, ClassificationSelection.id)
        ).all() if result_ids else []
        selections_by_result: dict[int, list[ClassificationSelection]] = {}
        active_selection_by_result: dict[int, ClassificationSelection] = {}
        for sel in selections:
            selections_by_result.setdefault(sel.classification_result_id, []).append(sel)
        for r_id in result_ids:
            r_sels = selections_by_result.get(r_id, [])
            superseded = {s.supersedes_selection_id for s in r_sels if s.supersedes_selection_id is not None}
            active = [s for s in r_sels if s.id not in superseded]
            if active:
                active_selection_by_result[r_id] = active[-1]
            elif r_sels:
                active_selection_by_result[r_id] = r_sels[-1]

        approval_events = session.scalars(
            select(ApprovalEvent)
            .where(ApprovalEvent.classification_run_id.in_(run_ids))
            .order_by(ApprovalEvent.created_at, ApprovalEvent.id)
        ).all() if run_ids else []

        makers = {item.id: item for item in session.scalars(select(Manufacturer)).all()}
        rules = {item.id: item for item in session.scalars(select(RuleSet)).all()}
        products_by_id = {p.id: p for p in products}
        corrections_query = select(Correction).where(Correction.certificate_id.in_(certificate_ids))
        if restrict_heats:
            corrections_query = corrections_query.where(Correction.replacement_observation_id.in_([o.id for o in observations]))
        corrections = session.scalars(corrections_query.order_by(Correction.id)).all()

        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in self.sheet_names:
            workbook.create_sheet(name)

        notice = self.settings.demo_notice if official else f"PRELIMINAR — PENDIENTE DE APROBACIÓN — {self.settings.demo_notice}"

        # 1. Resumen Sheet
        summary_rows = [
            ["Aviso Legal", notice],
            ["Generado UTC", datetime.now(timezone.utc).isoformat()],
            ["Tipo de Reporte", "Oficial (sólo expedientes aprobados)" if official else "Preliminar (borrador de trabajo)"],
            ["Actas incluidas", len(certificates)],
            ["Coladas incluidas", len(heats)],
            ["Rollos incluidos", len(products)],
            ["Clasificaciones incluidas", len(results)],
            ["Selecciones registradas", len(selections)],
            ["Ejecuciones de clasificación", ", ".join(str(run.id) for run in runs) or "Ninguna"],
        ] + [[f"Hoja {name}", f"Ir a hoja {name}"] for name in self.sheet_names if name != "Resumen"]
        self._write_sheet(workbook["Resumen"], ["Concepto", "Detalle"], summary_rows, "ResumenTable", notice)

        # Internal hyperlinks on Resumen
        for row in range(4, workbook["Resumen"].max_row + 1):
            label = workbook["Resumen"].cell(row, 1).value
            if isinstance(label, str) and label.startswith("Hoja "):
                target = label.removeprefix("Hoja ")
                link_cell = workbook["Resumen"].cell(row, 2)
                link_cell.hyperlink = f"#'{target}'!A1"
                link_cell.font = Font(color="0563C1", underline="single", bold=True)

        # 2. Actas Sheet
        self._write_sheet(workbook["Actas"],
            ["certificate_id", "document_id", "número", "fabricante", "fecha_acta", "fecha_carga", "estado", "revisión"],
            [[c.id, c.document_id, c.certificate_no, makers.get(c.manufacturer_id).name if c.manufacturer_id in makers else None,
              c.certificate_date, c.uploaded_at, c.approval_status, c.revision_number] for c in certificates],
            "ActasTable", notice)

        # 3. Coladas Sheet
        self._write_sheet(workbook["Coladas"],
            ["heat_id", "certificate_id", "colada", "norma", "grado", "propiedades"],
            [[h.id, h.certificate_id, h.heat_no, h.standard, h.grade, str(h.properties_json)] for h in heats],
            "ColadasTable", notice)

        # 4. Rollos Sheet
        self._write_sheet(workbook["Rollos"],
            ["product_id", "certificate_id", "heat_id", "identificador", "etiqueta", "tipo", "forma", "enrollado", "laminado", "ancho_mm", "espesor_mm", "peso_kg"],
            [[p.id, p.certificate_id, p.heat_id, p.product_identifier, p.label_no, p.product_type, p.form,
              p.coiled, p.rolling, p.width_mm, p.thickness_mm, p.weight_kg] for p in products],
            "RollosTable", notice)

        # 5. Composición Sheet
        self._write_sheet(workbook["Composición"],
            ["composition_id", "heat_id", "product_id", "elemento", "valor_original", "porcentaje", "heredado", "etiqueta_fuente"],
            [[c.id, c.heat_id, c.product_id, c.element, str(c.raw_value_json) if c.raw_value_json is not None else None,
              c.percentage, c.inherited, c.source_label] for c in chemistry],
            "ComposicionTable", notice)

        # 7. Evidencia Sheet (Built before Clasificación to map product row indices)
        evidencia_rows = []
        for row_idx, o in enumerate(observations, start=4):
            if o.product_id and o.product_id not in obs_by_product:
                obs_by_product[o.product_id] = row_idx
            regla_asociada = "; ".join(rule_by_observation.get(o.id, []))
            evidencia_rows.append([
                o.id, o.certificate_id, o.heat_id, o.product_id, o.field_path,
                str(o.raw_value_json) if o.raw_value_json is not None else None,
                str(o.normalized_value_json) if o.normalized_value_json is not None else None,
                o.unit, o.confidence, o.page_number, str(o.bbox_json) if o.bbox_json else None,
                o.source_text, regla_asociada, o.is_current,
            ])
        self._write_sheet(workbook["Evidencia"],
            ["observation_id", "certificate_id", "heat_id", "product_id", "campo", "original", "normalizado", "unidad", "confianza", "página", "coordenadas", "texto_fuente", "regla_factor", "vigente"],
            evidencia_rows, "EvidenciaTable", notice)

        # 6. Clasificación Sheet
        run_by_id = {run.id: run for run in runs}
        clasificacion_rows = []
        evidence_link_cells: list[tuple[int, int]] = []  # (row_index, target_evidence_row)

        for result_idx, r in enumerate(results, start=4):
            res_cands = candidates_by_result.get(r.id, [])
            cands_by_rank = {c.rank: c for c in res_cands}

            # Candidate 1, 2, 3 are ordered by rank in res_cands.

            # Selection resolution
            active_sel = active_selection_by_result.get(r.id)
            if active_sel:
                active_cand = next((c for c in res_cands if c.id == active_sel.candidate_id), None)
                chosen_code = f"{active_cand.fraction}-{active_cand.nico}" if active_cand else (f"{r.fraction}-{r.nico}" if r.fraction and r.nico else "Seleccionado")
                chosen_by = active_sel.person_name
                chosen_reason = active_sel.reason
                chosen_date = active_sel.created_at.isoformat()
            elif r.outcome == "classified" and r.fraction and r.nico:
                chosen_code = f"{r.fraction}-{r.nico}"
                chosen_by = "Motor determinista"
                chosen_reason = ("Resolución unívoca aprobada"
                                 if run_by_id[r.classification_run_id].approval_status == "approved"
                                 else "Resolución unívoca pendiente de aprobación")
                chosen_date = r.created_at.isoformat()
            else:
                chosen_code = "Pendiente de selección"
                chosen_by = "—"
                chosen_reason = "Pendiente de revisión documental"
                chosen_date = "—"

            # Factores explicativos
            active_cand_id = active_sel.candidate_id if active_sel else (cands_by_rank[1].id if 1 in cands_by_rank else None)
            res_factors = factors_by_candidate.get(active_cand_id, []) if active_cand_id else []
            factores_str = "; ".join(f"{f.rule_code}: {f.outcome}" for f in res_factors[:4]) if res_factors else "Reglas de partida y subpartida"

            prod = products_by_id.get(r.product_id)
            prod_ident = prod.product_identifier if prod else str(r.product_id)

            alternatives = [c for c in res_cands if c.id != active_cand_id]
            cand2_code = f"{alternatives[0].fraction}-{alternatives[0].nico}" if alternatives else "—"
            cand3_code = f"{alternatives[1].fraction}-{alternatives[1].nico}" if len(alternatives) > 1 else "—"
            target_ev_row = obs_by_product.get(r.product_id)
            if target_ev_row is None:
                target_ev_row = next((i for i, o in enumerate(observations, start=4)
                                      if prod and o.product_id is None
                                      and o.heat_id == prod.heat_id and o.heat_id is not None), None)
            clasificacion_rows.append([
                r.id, r.classification_run_id, r.product_id, prod_ident,
                chosen_code, cand2_code, cand3_code,
                chosen_by, chosen_reason, chosen_date,
                factores_str, "Ver evidencia" if target_ev_row is not None else "Sin evidencia",
                r.outcome,
                ", ".join(r.details_json.get("missing_fields") or []) or "Ninguno",
                run_by_id[r.classification_run_id].approval_status,
                rules.get(run_by_id[r.classification_run_id].rule_set_id).version if run_by_id[r.classification_run_id].rule_set_id in rules else "—",
            ])

            if target_ev_row is not None:
                evidence_link_cells.append((result_idx, target_ev_row))

        self._write_sheet(workbook["Clasificación"],
            ["result_id", "run_id", "product_id", "rollo_identificador", "candidato_elegido", "alternativa_2", "alternativa_3",
             "seleccionado_por", "motivo_selección", "fecha_selección", "factores_clave", "evidencia_enlace",
             "resultado", "faltantes", "estado_ejecución", "reglas_versión"],
            clasificacion_rows, "ClasificacionTable", notice)

        # Apply internal evidence links in Clasificación sheet
        for clasif_row, target_ev_row in evidence_link_cells:
            link_cell = workbook["Clasificación"].cell(clasif_row, 12)
            link_cell.hyperlink = f"#'Evidencia'!A{target_ev_row}"
            link_cell.font = Font(color="0563C1", underline="single", bold=True)

        # 8. Auditoría Sheet (Unified timeline of selections, approvals, and corrections)
        auditoria_rows = []

        # Selections timeline
        for sel in selections:
            sel_cand = next((c for c in candidates if c.id == sel.candidate_id), None)
            cand_repr = f"{sel_cand.fraction}-{sel_cand.nico}" if sel_cand else f"candidato-{sel.candidate_id}"
            res = next((r for r in results if r.id == sel.classification_result_id), None)
            is_active = active_selection_by_result.get(sel.classification_result_id) == sel
            auditoria_rows.append([
                "SELECCION_CANDIDATO", sel.id,
                run_by_id[res.classification_run_id].certificate_id if res else None,
                res.classification_run_id if res else None,
                res.product_id if res else None,
                f"Seleccionó {cand_repr} (Reemplaza selección #{sel.supersedes_selection_id})" if sel.supersedes_selection_id else f"Seleccionó {cand_repr}",
                sel.person_name, sel.reason, sel.workstation_name, sel.created_at.isoformat(),
                "Vigente" if is_active else "Reemplazada",
            ])

        # Approval events timeline
        latest_approval = {event.classification_run_id: event.id for event in approval_events}
        for app in approval_events:
            auditoria_rows.append([
                "APROBACION_ESTADO", app.id,
                run_by_id[app.classification_run_id].certificate_id if app.classification_run_id in run_by_id else None,
                app.classification_run_id, None,
                f"Cambio de estado: {app.from_status} -> {app.to_status}",
                app.person_name, app.reason, app.workstation_name, app.created_at.isoformat(),
                "Vigente" if latest_approval[app.classification_run_id] == app.id else "Reemplazada",
            ])

        # Corrections timeline
        observations_by_id = {o.id: o for o in observations}
        for cor in corrections:
            replacement = observations_by_id.get(cor.replacement_observation_id)
            auditoria_rows.append([
                "CORRECCION_DATO", cor.id, cor.certificate_id, None, None,
                f"Corrección obs {cor.previous_observation_id} -> {cor.replacement_observation_id}",
                cor.person_name, cor.reason, cor.workstation_name, cor.created_at.isoformat(),
                "Vigente" if replacement and replacement.is_current else "Reemplazada",
            ])

        auditoria_rows.sort(key=lambda item: str(item[9]))

        self._write_sheet(workbook["Auditoría"],
            ["tipo_evento", "id_evento", "certificate_id", "run_id", "product_id", "detalle", "persona", "motivo", "estación", "fecha_utc", "vigencia"],
            auditoria_rows, "AuditoriaTable", notice)

        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        self._validate(destination, len(certificates), len(heats), len(products))
        return destination

    @staticmethod
    def _write_sheet(sheet, headers: list[str], rows: Iterable[list[Any]], table_name: str, notice: str) -> None:
        sheet["A1"] = notice
        sheet["A1"].font = Font(bold=True, color="9C0006")
        sheet.append([])
        sheet.append(headers)
        for row_index, row in enumerate(rows, start=4):
            clean_row = [
                cell_val.astimezone(timezone.utc).replace(tzinfo=None)
                if isinstance(cell_val, datetime) and cell_val.tzinfo is not None
                else cell_val
                for cell_val in row
            ]
            sheet.append(clean_row)
            # External text must remain text, even when Excel recognizes a formula.
            for column_index, value in enumerate(clean_row, start=1):
                if isinstance(value, str):
                    sheet.cell(row_index, column_index).data_type = "s"
        for cell in sheet[3]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{sheet.cell(3, len(headers)).coordinate}"
        if sheet.max_row >= 4:
            table = Table(displayName=table_name, ref=f"A3:{sheet.cell(sheet.max_row, len(headers)).coordinate}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)

    @staticmethod
    def _validate(path: Path, certificates: int, heats: int, products: int) -> None:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            expected_sheets = {
                "Resumen", "Actas", "Coladas", "Rollos", "Composición",
                "Clasificación", "Evidencia", "Auditoría",
            }
            if set(workbook.sheetnames) != expected_sheets:
                raise RuntimeError(f"El libro no contiene exactamente las 8 hojas requeridas: {workbook.sheetnames}")
            if len(list(workbook["Actas"].iter_rows(min_row=4, values_only=True))) != certificates:
                raise RuntimeError("El conteo de actas exportadas no coincide")
            if len(list(workbook["Coladas"].iter_rows(min_row=4, values_only=True))) != heats:
                raise RuntimeError("El conteo de coladas exportadas no coincide")
            if len(list(workbook["Rollos"].iter_rows(min_row=4, values_only=True))) != products:
                raise RuntimeError("El conteo de rollos exportados no coincide")
        finally:
            workbook.close()
