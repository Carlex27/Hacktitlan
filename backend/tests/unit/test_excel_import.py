from io import BytesIO
from zipfile import ZipFile
from types import SimpleNamespace

from openpyxl import Workbook
import pytest
from fastapi.testclient import TestClient

from backend.app.application.certificate_extraction import CertificateExtractionService
from backend.app.certificate_parser.excel_certificate import extract_excel
from backend.app.domain.errors import ApplicationError
from backend.app.infrastructure.files import FileStorage
from backend.app.api.app import create_app, get_session
from backend.app.infrastructure.database.models import Document, EvidenceLink, MillCertificate, Observation, StoredFile


def mill_book() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mill"
    sheet.append(["MILL NO", "COIL NO", "PRODUCTION NAME", "WIDTH", "THICK", "LENGTH",
                  "COIL WEIGHT(MT)", "溶鋼分析・Ｃ・実績値", "溶鋼分析・Ｃ・有効桁",
                  "溶鋼分析・成分・元素名-01", "溶鋼分析・成分・実績値-01",
                  "溶鋼分析・成分・有効桁-01", "FRACCION", "NICO"])
    sheet.append(["CERT-1", "COIL-1", "HOT ROLLED STEEL SHEET IN COIL", 1000, 1.99, 0,
                  20.07, .05, 2, "TI", .05, 2, "72253099", "01"])
    sheet.append(["CERT-2", None, "STEEL SHEET", 900, .9, None,
                  None, None, None, None, 0, 2, '=IF(A1="",0,1)', None])
    sheet.append(["CERT-3", None, "STEEL SHEET", 900, .9, None,
                  None, 0, 2, "NB", 0, 2, None, None])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_excel_upload_and_normalization_preserve_zero_missing_and_source(tmp_path):
    storage = FileStorage(tmp_path, 1_000_000)
    content = mill_book()
    artifact = storage.store_document(BytesIO(content), "mill.xlsx")
    assert artifact.media_type.endswith("spreadsheetml.sheet")
    path = storage.resolve(artifact.relative_path)
    assert path.read_bytes() == content
    assert storage.store_document(BytesIO(content), "renamed.xlsx").duplicate
    result = extract_excel(path)
    first, second, third = result["certificate"]["products"]
    assert first["composition_pct"] == {"C": .05, "Ti": .05}
    assert first["weight_kg"] == 20070
    assert first["length_m"] is None
    assert first["coiled"] is True
    assert first["rolling"] == "hot"
    assert first["observations"]["composition_pct"]["Ti"]["source_label"] == "'Mill'!K2"
    assert second["composition_pct"] == {"C": None}
    assert second["coiled"] is None
    assert third["composition_pct"] == {"C": 0, "Nb": 0}
    assert second["observations"]["source_fraction"]["raw_value"].startswith("=IF")
    assert result["status"] == "needs_review"
    assert result["document"]["page_count"] is None


def test_general_inverted_headers_withhold_chemistry(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["MILL NO", "PRODUCTION NAME", "HEAT No.", "THICK (mm)", "WIDTH (mm)", "LENGTH",
                  "溶鋼分析・Ｃ・実績値", "溶鋼分析・Ｃ・有効桁"])
    sheet.append(["CERT-1", "COLD-ROLLED STEEL STRIP", "HEAT-1", .9, 991, "C", 4, .0013])
    path = tmp_path / "general.xlsx"
    workbook.save(path)
    result = extract_excel(path)
    product = result["certificate"]["products"][0]
    assert product["composition_pct"]["C"] is None
    assert product["thickness_mm"] == .9
    assert product["heat_no"] == "HEAT-1"
    assert product["coiled"] is True
    assert result["document"]["ingestion"]["sheets"][0]["rows"][1]["cells"]["H2"] == .0013
    assert any("invertidos" in reason for reason in result["reasons"])


@pytest.mark.parametrize("sheet_name", ["KIMITSU(MILL CERT)", "Unknown mill"])
def test_kimitsu_inverted_pairs_preserve_percentages_and_cell_evidence(tmp_path, sheet_name):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(["MILL NO", "PRODUCTION NAME", "HEAT No.", "THICK (mm)", "WIDTH (mm)", "LENGTH",
                  *[f"溶鋼分析・{element}・{suffix}" for element in ("Ｃ", "Ｓｉ", "Ｍｎ", "Ｐ", "Ｓ")
                    for suffix in ("実績値", "有効桁")],
                  "溶鋼分析・ Al・有効桁", "溶鋼分析・ Al・有効桁",
                  "溶鋼分析・ Ti・有効桁", "溶鋼分析・ Ti・有効桁"])
    sheet.append(["E02511200001", "COLD-ROLLED STEEL STRIP", "25BH2B7550200", .9, 991, "C",
                  4, .0013, None, 0, 2, .12, 3, .011, 3, .008, 3, .031, 3, .068])
    sheet.append(["CERT-2", "STEEL STRIP", "HEAT-2", .9, 991, "C",
                  .5, .0013, None, None, 2, .12])
    path = tmp_path / "kimitsu.xlsx"
    workbook.save(path)
    result = extract_excel(path)
    first, second = result["certificate"]["products"]
    assert first["heat_no"] == "25BH2B7550200"
    assert result["status"] == "needs_review"
    if sheet_name == "KIMITSU(MILL CERT)":
        assert first["composition_pct"] == {
            "C": .0013, "Si": 0, "Mn": .12, "P": .011, "S": .008, "Al_total": .031, "Ti": .068,
        }
        detail = first["observations"]["composition_pct"]["C"]
        assert detail["source_label"] == "'KIMITSU(MILL CERT)'!H2"
        assert detail["raw_value"] == detail["normalized_value"] == .0013
        assert second["composition_pct"]["C"] is None
        assert second["composition_pct"]["Si"] is None
    else:
        assert all(value is None for value in first["composition_pct"].values())


def test_cast_number_identifies_heat_without_replacing_coil_number(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CAST NO", "COIL NO", "WIDTH", "THICK", "C"])
    sheet.append(["G82765", "603214670", 1183, 1.99, .05])
    path = tmp_path / "cast.xlsx"
    workbook.save(path)
    product = extract_excel(path)["certificate"]["products"][0]
    assert product["heat_no"] == "G82765"
    assert product["label_no"] == "603214670"
    assert product["observations"]["heat_no"]["source_label"] == "'Sheet'!A2"


def test_support_letter_preserves_merged_bounds_without_fake_products(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["SPEC", "Alloy  or Non-Alloy", "Yield Point\n(N/mm2)"])
    sheet.append(["SPCC", "non-alloy", "＜355"])
    sheet.append(["SPCD", "non-alloy", None])
    sheet.merge_cells("C2:C3")
    sheet.append(["OTHER", "alloy", "≧355"])
    path = tmp_path / "support.xlsx"
    workbook.save(path)
    result = extract_excel(path)
    assert result["certificate"]["products"] == []
    records = result["document"]["ingestion"]["specification_records"]
    assert records[0]["yield_strength"]["normalized_value"] == {"operator": "<", "value": 355, "unit": "MPa"}
    assert records[1]["yield_strength"]["inherited"] is True
    assert records[1]["yield_strength"]["source_label"].endswith("!C2")
    assert records[2]["yield_strength"]["normalized_value"]["operator"] == ">="


def test_excel_bypasses_pdf_reader_and_honors_cancellation(tmp_path):
    class NoPdfReader:
        def read(self, *_args, **_kwargs):
            raise AssertionError("Excel must not enter OCR")
    path = tmp_path / "mill.xlsx"
    path.write_bytes(mill_book())
    service = CertificateExtractionService(NoPdfReader())
    assert service.analyze_document(path)["status"] == "needs_review"
    with pytest.raises(Exception, match="cancelada"):
        service.analyze_document(path, cancel_check=lambda: True)


def test_unknown_sheet_and_formula_descriptions_remain_unknown(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["MILL NO", "PRODUCTION NAME", "WIDTH", "THICK", "C"])
    sheet.append(['="CERT-1"', '="HOT ROLLED STEEL SHEET IN COIL"', 1000, 2, "=1/10"])
    unknown = workbook.create_sheet("Other")
    unknown.append(["Keep original", "=SUM(1,2)"])
    workbook.save(path)
    result = extract_excel(path)
    product = result["certificate"]["products"][0]
    assert product["source_certificate_no"] is None
    assert product["form"] is None
    assert product["rolling"] is None
    assert product["coiled"] is None
    assert product["composition_pct"]["C"] is None
    assert result["document"]["ingestion"]["sheets"][1]["rows"][0]["cells"]["B1"] == "=SUM(1,2)"


def test_oversized_merged_range_rejected_before_loading(tmp_path):
    content = mill_book()
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(output, "w") as changed:
        for entry in original.infolist():
            value = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                value = value.replace(b"</worksheet>", b'<mergeCells><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>')
            changed.writestr(entry, value)
    with pytest.raises(ApplicationError) as error:
        FileStorage(tmp_path, 1_000_000).store_document(BytesIO(output.getvalue()), "huge.xlsx")
    assert error.value.status_code == 413


def test_excel_evidence_returns_original_file_instead_of_pdf_page():
    records = {
        EvidenceLink: SimpleNamespace(id=1, source_type="observation", observation_id=2,
                                      decision_step_id=3, candidate_factor_id=None, field_path="composition_pct.Ti"),
        Observation: SimpleNamespace(id=2, certificate_id=4, raw_value_json=.05, normalized_value_json=.05,
                                    unit="%", confidence=1, source_text="'Mill'!K2: 0.05", page_number=None, bbox_json=None),
        MillCertificate: SimpleNamespace(id=4, document_id=5),
        Document: SimpleNamespace(id=5, stored_file_id=6),
        StoredFile: SimpleNamespace(id=6, original_name="mill.xlsx"),
    }
    class EvidenceSession:
        def get(self, model, _identifier):
            return records[model]
    app = create_app()
    app.dependency_overrides[get_session] = lambda: EvidenceSession()
    response = TestClient(app).get("/api/v1/evidence/1")
    assert response.status_code == 200
    assert response.json()["data"]["focus"]["fallback"] == "original_file"
    assert response.json()["data"]["observation"]["source_text"] == "'Mill'!K2: 0.05"


@pytest.mark.parametrize("content", [b"%PDF-1.7", b"PK fake", b""])
def test_invalid_xlsx_rejected_and_temp_removed(tmp_path, content):
    storage = FileStorage(tmp_path, 1_000_000)
    with pytest.raises(ApplicationError) as error:
        storage.store_document(BytesIO(content), "bad.xlsx")
    assert error.value.code == "invalid_xlsx"
    assert list(storage.temp_root.iterdir()) == []


def test_size_limits_macros_and_legacy_formats(tmp_path):
    with pytest.raises(ApplicationError) as error:
        FileStorage(tmp_path, 5).store_document(BytesIO(mill_book()), "large.xlsx")
    assert error.value.status_code == 413
    with pytest.raises(ApplicationError) as error:
        FileStorage(tmp_path, 1_000_000).store_document(BytesIO(b"xls"), "legacy.xls")
    assert error.value.code == "invalid_file_type"
    output = BytesIO(mill_book())
    with ZipFile(output, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(ApplicationError) as error:
        FileStorage(tmp_path, 1_000_000).store_document(BytesIO(output.getvalue()), "macro.xlsx")
    assert error.value.code == "invalid_xlsx"
