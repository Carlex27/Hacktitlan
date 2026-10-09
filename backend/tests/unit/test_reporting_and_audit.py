from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
import base64
import json
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from backend.app.api.app import create_app, get_session
from backend.app.api.schemas import ExportRequest
from backend.app.application.classification_service import ClassificationService
from backend.app.application.document_service import (
    DocumentService,
    decode_cursor,
    encode_cursor,
    resolve_period_dates,
)
from backend.app.config import Settings
from backend.app.domain.errors import ApplicationError
from backend.app.infrastructure.database.models import (
    ApprovalEvent,
    CandidateFactor,
    ChemicalComposition,
    ClassificationCandidate,
    ClassificationResult,
    ClassificationRun,
    ClassificationSelection,
    Correction,
    Document,
    EvidenceLink,
    Export,
    Heat,
    Job,
    Manufacturer,
    MillCertificate,
    Observation,
    Product,
    RuleSet,
    StoredFile,
)
from backend.app.infrastructure.files import FileStorage
from backend.app.reporting.excel_export import ExcelExportService


# ---------------------------------------------------------------------------
# Unit tests: pure logic and schema validation
# ---------------------------------------------------------------------------

def test_resolve_period_dates_presets():
    today = date(2026, 10, 8)
    assert resolve_period_dates("today", today=today) == (today, today)
    assert resolve_period_dates("day", today=today) == (today, today)
    assert resolve_period_dates("week", today=today) == (today - timedelta(days=7), today)
    assert resolve_period_dates("month", today=today) == (today - timedelta(days=30), today)
    assert resolve_period_dates(None, today=today) == (None, None)
    assert resolve_period_dates("all", today=today) == (None, None)
    assert resolve_period_dates("custom", today=today) == (None, None)


def test_cursor_roundtrip_and_errors():
    d = date(2026, 10, 8)
    encoded = encode_cursor(d, 42)
    assert decode_cursor(encoded) == (d, 42)

    with pytest.raises(ApplicationError) as exc_info:
        decode_cursor("not_a_valid_cursor")
    assert exc_info.value.code == "invalid_cursor"


@pytest.mark.parametrize("payload", [
    ["2026-10-08", True], ["2026-10-08", 1.5], ["2026-10-08", "42"],
    ["2026-10-08", 0], ["2026-10-08", -1], ["2026-10-08", 2**63], ["2026-10-08", 1, "extra"],
    {"0": "2026-10-08", "1": 1},
])
def test_cursor_rejects_invalid_payload(payload):
    encoded = base64.urlsafe_b64encode(json.dumps(payload).encode()).decode()
    with pytest.raises(ApplicationError, match="cursor"):
        decode_cursor(encoded)


@pytest.mark.parametrize("scope", ["certificate_ids", "heat_ids", "classification_run_ids"])
def test_export_request_requires_positive_ids(scope):
    with pytest.raises(ValueError):
        ExportRequest(person_name="Auditor", **{scope: [0]})


def test_excel_external_text_is_not_a_formula(tmp_path):
    workbook = openpyxl.Workbook()
    text = '=HYPERLINK("https://example.test", "dato")'
    ExcelExportService._write_sheet(workbook.active, ["texto", "fecha"],
        [[text, datetime(2026, 10, 8, 12, tzinfo=timezone(timedelta(hours=-6)))]], "TextTable", "Aviso")
    path = tmp_path / "text.xlsx"
    workbook.save(path)
    loaded = openpyxl.load_workbook(path)
    assert loaded.active["A4"].value == text
    assert loaded.active["A4"].data_type == "s"
    assert loaded.active["B4"].value == datetime(2026, 10, 8, 18)
    loaded.close()


def test_export_request_validation():
    with pytest.raises(ValueError):
        ExportRequest(person_name="Test User")

    req = ExportRequest(
        person_name="Auditor",
        workstation_name="WS-01",
        certificate_ids=[1, 2],
        official=True,
        filters={"period": "month", "heat_no": "HEAT-99"},
    )
    assert req.certificate_ids == [1, 2]
    assert req.official is True
    assert req.filters == {"period": "month", "heat_no": "HEAT-99"}


# ---------------------------------------------------------------------------
# Database-backed tests: query filters, pagination, excel export & API
# ---------------------------------------------------------------------------

@pytest.fixture()
def db_session():
    db_url = Settings().test_database_url
    if not db_url:
        pytest.skip("Requires HACKTITLAN_TEST_DATABASE_URL")
    assert db_url.rsplit("/", 1)[-1].split("?", 1)[0] == "hacktitlan_test"
    engine = create_engine(db_url)
    connection = engine.connect()
    transaction = connection.begin()
    session = Session(bind=connection)
    try:
        yield session, db_url
    finally:
        session.close()
        if transaction.is_active:
            transaction.rollback()
        connection.close()
        engine.dispose()


def _seed_test_document_and_cert(session: Session, cert_no: str, cert_date: date, status: str = "draft"):
    token = uuid4().hex
    stored = StoredFile(
        sha256=token.ljust(64, "0"),
        original_name=f"{cert_no}.pdf",
        storage_name=f"{token}.pdf",
        relative_path=f"test/{token}.pdf",
        media_type="application/pdf",
        size_bytes=1024,
    )
    session.add(stored)
    session.flush()

    mfr = Manufacturer(name=f"Maker-{token[:6]}", normalized_name=f"maker-{token[:6]}")
    session.add(mfr)
    session.flush()

    doc = Document(
        stored_file_id=stored.id,
        processing_status="succeeded",
    )
    session.add(doc)
    session.flush()

    cert = MillCertificate(
        document_id=doc.id,
        manufacturer_id=mfr.id,
        certificate_no=cert_no,
        certificate_date=cert_date,
        approval_status=status,
        demo_notice="DEMOSTRACIÓN — SIN VALIDEZ ADUANERA",
    )
    session.add(cert)
    session.flush()
    return doc, cert, mfr


def test_document_service_list_certificates_filters(db_session, tmp_path):
    session, db_url = db_session
    settings = Settings(database_url=db_url, storage_root=tmp_path / "storage")
    service = DocumentService(settings, FileStorage(settings.storage_root, 1024))
    today = date.today()

    _, cert1, mfr1 = _seed_test_document_and_cert(session, "CERT-FILTER-01", today, status="approved")
    _, cert2, _ = _seed_test_document_and_cert(session, "CERT-FILTER-02", today - timedelta(days=20), status="needs_review")

    heat1 = Heat(certificate_id=cert1.id, heat_no="HEAT-ALPHA-01")
    session.add(heat1)
    session.flush()
    prod1 = Product(certificate_id=cert1.id, heat_id=heat1.id, product_identifier="COIL-ALPHA-01")
    session.add(prod1)
    session.flush()

    rs = session.scalar(select(RuleSet).order_by(RuleSet.id).limit(1))
    assert rs is not None
    run1 = ClassificationRun(
        certificate_id=cert1.id,
        rule_set_id=rs.id,
        approval_status="approved",
        input_snapshot_json={},
        demo_notice="DEMOSTRACIÓN — SIN VALIDEZ ADUANERA",
    )
    session.add(run1)
    session.flush()
    res1 = ClassificationResult(
        classification_run_id=run1.id,
        product_id=prod1.id,
        outcome="classified",
        fraction="72091504",
        nico="01",
    )
    session.add(res1)
    session.flush()

    # 1. Filter by certificate_no
    items, _ = service.list_certificates(session, limit=50, certificate_no="CERT-FILTER-01")
    assert any(c.id == cert1.id for c, _ in items)
    assert not any(c.id == cert2.id for c, _ in items)

    # 2. Filter by heat_no
    items, _ = service.list_certificates(session, limit=50, heat_no="ALPHA-01")
    assert any(c.id == cert1.id for c, _ in items)

    # 3. Filter by product_identifier
    items, _ = service.list_certificates(session, limit=50, product_identifier="COIL-ALPHA")
    assert any(c.id == cert1.id for c, _ in items)

    # 4. Filter by fraction and nico
    items, _ = service.list_certificates(session, limit=50, fraction="72091504", nico="01")
    assert any(c.id == cert1.id for c, _ in items)

    other_product = Product(certificate_id=cert1.id, heat_id=heat1.id, product_identifier="OTHER")
    session.add(other_product)
    session.flush()
    session.add(ClassificationResult(classification_run_id=run1.id, product_id=other_product.id,
                                    outcome="classified", fraction="72099999", nico="02"))
    session.flush()
    items, _ = service.list_certificates(session, limit=50, fraction="72091504", nico="02")
    assert cert1.id not in [c.id for c, _ in items]
    with pytest.raises(ApplicationError, match="fecha inicial"):
        service.list_certificates(session, limit=50, date_from=today, date_to=today - timedelta(days=1))

    # 5. Filter by period="week" vs "month"
    items_week, _ = service.list_certificates(session, limit=50, period="week")
    assert any(c.id == cert1.id for c, _ in items_week)
    assert not any(c.id == cert2.id for c, _ in items_week)

    items_month, _ = service.list_certificates(session, limit=50, period="month")
    assert any(c.id == cert1.id for c, _ in items_month)
    assert any(c.id == cert2.id for c, _ in items_month)

    # 6. Filter by approval_status
    items_appr, _ = service.list_certificates(session, limit=50, approval_status="approved")
    assert any(c.id == cert1.id for c, _ in items_appr)
    assert not any(c.id == cert2.id for c, _ in items_appr)


def test_history_filters_match_visible_documents_and_current_rolls(db_session, tmp_path):
    session, db_url = db_session
    settings = Settings(database_url=db_url, storage_root=tmp_path / "storage", backup_root=tmp_path / "backups")
    doc, cert, maker = _seed_test_document_and_cert(session, "FILTER-CURRENT", date(2026, 6, 12), "needs_review")
    _, unknown, _ = _seed_test_document_and_cert(session, "FILTER-UNKNOWN", None)
    doc.stored_file.original_name = "MOLINOS GENERAL.xlsx"
    heat = Heat(certificate_id=cert.id, heat_no="HEAT-ONE")
    other_heat = Heat(certificate_id=cert.id, heat_no="HEAT-TWO")
    session.add_all([heat, other_heat])
    session.flush()
    roll = Product(certificate_id=cert.id, heat_id=heat.id, product_identifier="ROLL_ONE")
    other_roll = Product(certificate_id=cert.id, heat_id=other_heat.id, product_identifier="ROLL-TWO")
    session.add_all([roll, other_roll])
    session.flush()
    rule_set = session.scalar(select(RuleSet).order_by(RuleSet.id).limit(1))
    old_run = ClassificationRun(certificate_id=cert.id, rule_set_id=rule_set.id,
        input_snapshot_json={}, demo_notice="Test", created_at=datetime(2026, 1, 1, tzinfo=timezone.utc))
    new_run = ClassificationRun(certificate_id=cert.id, rule_set_id=rule_set.id,
        input_snapshot_json={}, demo_notice="Test", created_at=datetime(2026, 1, 2, tzinfo=timezone.utc))
    session.add_all([old_run, new_run])
    session.flush()
    session.add_all([
        ClassificationResult(classification_run_id=old_run.id, product_id=roll.id,
            outcome="classified", fraction="72091504", nico="01"),
        ClassificationResult(classification_run_id=new_run.id, product_id=roll.id,
            outcome="classified", fraction="72104999", nico="02"),
        ClassificationResult(classification_run_id=new_run.id, product_id=other_roll.id,
            outcome="classified", fraction="72091504", nico="03"),
    ])
    session.flush()
    app = create_app(settings)
    def get_test_session():
        yield session
    app.dependency_overrides[get_session] = get_test_session
    with TestClient(app) as client:
        def ids(**params):
            response = client.get("/api/v1/certificates", params=params)
            assert response.status_code == 200, response.text
            return [item["id"] for item in response.json()["data"]]

        for search in (" filter-current ", " molinos general ", str(cert.id), f"Acta #{cert.id}"):
            assert cert.id in ids(certificate_no=search)
        assert cert.id in ids(manufacturer=f" {maker.name.lower()} ")
        assert cert.id in ids(heat_no=" heat-one ", product_identifier=" roll_one ", fraction="7210.49.99", nico=" 02 ", approval_status="needs_review")
        assert cert.id not in ids(heat_no="HEAT-TWO", product_identifier="ROLL_ONE")
        assert cert.id not in ids(product_identifier="ROLL_ONE", fraction="72091504")
        assert cert.id not in ids(fraction="72104999", nico="03")
        assert cert.id not in ids(product_identifier="ROLL%ONE")
        assert cert.id not in ids(certificate_no="%")
        assert cert.id in ids(certificate_no="FILTER-CURRENT", date_from="2026-06-12", date_to="2026-06-12")
        assert cert.id not in ids(certificate_no="FILTER-CURRENT", date_from="2026-06-13")
        assert cert.id not in ids(certificate_no="FILTER-CURRENT", approval_status="approved")
        assert unknown.id not in ids(certificate_no="FILTER-UNKNOWN", date_from="2000-01-01")
        assert unknown.id in ids(certificate_no="FILTER-UNKNOWN", date_from="2000-01-01", date_basis="uploaded")
        assert cert.id in ids(manufacturer="   ")
        assert client.get("/api/v1/certificates", params={"date_from": "2026-06-13", "date_to": "2026-06-12"}).status_code == 400
        parameters = client.get("/openapi.json").json()["paths"]["/api/v1/certificates"]["get"]["parameters"]
        assert "nombre original" in next(p for p in parameters if p["name"] == "certificate_no")["description"]


@pytest.mark.parametrize("insert_day", [10, 1])
def test_cursor_pagination_stability_under_concurrent_insertions(db_session, tmp_path, insert_day):
    session, db_url = db_session
    settings = Settings(database_url=db_url, storage_root=tmp_path / "storage")
    service = DocumentService(settings, FileStorage(settings.storage_root, 1024))
    base_date = date(2026, 5, 1)
    prefix = f"PAGE-{uuid4().hex[:6]}"

    # Insert 4 certificates with distinct dates
    certs = []
    for i in range(4):
        _, cert, _ = _seed_test_document_and_cert(session, f"{prefix}-{i}", base_date + timedelta(days=i * 2))
        certs.append(cert)

    # Page 1 with limit=2 (most recent first: certs[3], certs[2])
    p1_items, next_cursor = service.list_certificates(session, limit=2, certificate_no=prefix)
    assert len(p1_items) == 2
    p1_ids = [cert.id for cert, _ in p1_items]
    assert next_cursor is not None

    # Simulate concurrent insertion of a brand new certificate between page fetches
    _, cert_concurrent, _ = _seed_test_document_and_cert(session, f"{prefix}-NEW", base_date + timedelta(days=insert_day))

    # Page 2 using cursor
    p2_items, last_cursor = service.list_certificates(session, limit=2, cursor=next_cursor, certificate_no=prefix)
    assert len(p2_items) == 2
    p2_ids = [cert.id for cert, _ in p2_items]

    # Keyset pagination invariant: NO duplicate IDs between page 1 and page 2
    assert set(p1_ids).isdisjoint(set(p2_ids))
    # The concurrent insertion must not alter the relative boundary of page 2
    rest, _ = service.list_certificates(session, limit=2, cursor=last_cursor, certificate_no=prefix) if last_cursor else ([], None)
    all_ids = p1_ids + p2_ids + [c.id for c, _ in rest]
    assert len(all_ids) == len(set(all_ids))
    assert [i for i in all_ids if i != cert_concurrent.id] == [c.id for c in reversed(certs)]
    assert (cert_concurrent.id in all_ids) == (insert_day == 1)


def test_classification_run_detail_exposes_current_selection_and_selections(db_session):
    session, db_url = db_session
    today = date.today()
    _, cert, _ = _seed_test_document_and_cert(session, "CERT-RUN-SEL", today, status="needs_review")

    prod = Product(certificate_id=cert.id, product_identifier="COIL-RUN-SEL")
    session.add(prod)
    session.flush()

    rs = session.scalar(select(RuleSet).order_by(RuleSet.id).limit(1))
    assert rs is not None
    run = ClassificationRun(
        certificate_id=cert.id,
        rule_set_id=rs.id,
        approval_status="needs_review",
        input_snapshot_json={},
        demo_notice="DEMOSTRACIÓN — SIN VALIDEZ ADUANERA",
    )
    session.add(run)
    session.flush()

    res = ClassificationResult(
        classification_run_id=run.id,
        product_id=prod.id,
        outcome="needs_review",
    )
    session.add(res)
    session.flush()

    c1 = ClassificationCandidate(
        classification_result_id=res.id,
        rank=1,
        fraction="72091504",
        nico="01",
        description="Candidate 1",
        support_level="fully_supported",
    )
    c2 = ClassificationCandidate(
        classification_result_id=res.id,
        rank=2,
        fraction="72091504",
        nico="02",
        description="Candidate 2",
        support_level="conditional",
    )
    session.add_all([c1, c2])
    session.flush()

    # Selection 1
    sel1 = ClassificationSelection(
        classification_result_id=res.id,
        candidate_id=c1.id,
        person_name="Analyst 1",
        reason="Initial choice",
        workstation_name="WS-1",
        created_at=datetime.now(timezone.utc) - timedelta(minutes=10),
    )
    session.add(sel1)
    session.flush()

    # Selection 2 (supersedes Selection 1)
    sel2 = ClassificationSelection(
        classification_result_id=res.id,
        candidate_id=c2.id,
        person_name="Supervisor 1",
        reason="Corrected standard verification",
        workstation_name="WS-2",
        supersedes_selection_id=sel1.id,
        created_at=datetime.now(timezone.utc),
    )
    session.add(sel2)
    session.flush()

    app = create_app(Settings(database_url=db_url))

    def _override_get_session():
        yield session

    app.dependency_overrides[get_session] = _override_get_session

    with TestClient(app) as client:
        res_detail = client.get(f"/api/v1/classification-runs/{run.id}")
        assert res_detail.status_code == 200
        detail = res_detail.json()["data"]
        result_entry = detail["results"][0]

        # Verify current_selection points to sel2
        assert result_entry["current_selection"] is not None
        assert result_entry["current_selection"]["id"] == sel2.id
        assert result_entry["current_selection"]["candidate_id"] == c2.id
        assert result_entry["current_selection"]["person_name"] == "Supervisor 1"
        assert result_entry["current_selection"]["reason"] == "Corrected standard verification"

        # Verify selections contains both sel1 and sel2
        assert len(result_entry["selections"]) == 2
        assert result_entry["selections"][0]["id"] == sel1.id
        assert result_entry["selections"][1]["id"] == sel2.id


def test_excel_export_service_generates_eight_sheets_and_auditable_content(db_session, tmp_path):
    session, db_url = db_session
    today = date.today()
    settings = Settings(database_url=db_url, storage_root=tmp_path / "storage")

    _, cert, mfr = _seed_test_document_and_cert(session, "CERT-AUDIT-EXCEL", today, status="approved")

    heat = Heat(certificate_id=cert.id, heat_no="HEAT-EXCEL-01", standard="ASTM A36", grade="Gr 36")
    session.add(heat)
    session.flush()

    prod = Product(
        certificate_id=cert.id,
        heat_id=heat.id,
        product_identifier="ROLLO-EXCEL-01",
        form="coil",
        coiled=True,
        thickness_mm=Decimal("3.00"),
        width_mm=Decimal("1200.00"),
        weight_kg=Decimal("15000.00"),
    )
    session.add(prod)
    session.flush()

    chem1 = ChemicalComposition(
        heat_id=heat.id,
        element="C",
        percentage=Decimal("0.15"),
        raw_value_json=0.15,
    )
    chem2 = ChemicalComposition(
        heat_id=heat.id,
        element="Mn",
        percentage=Decimal("0.80"),
        raw_value_json=0.80,
    )
    chem3 = ChemicalComposition(
        heat_id=heat.id,
        element="Si",
        percentage=Decimal("0.20"),
        raw_value_json=0.20,
    )
    session.add_all([chem1, chem2, chem3])
    session.flush()

    obs = Observation(
        certificate_id=cert.id,
        product_id=prod.id,
        field_path="thickness_mm",
        source_text="3.00 mm",
        raw_value_json="3.00 mm",
        normalized_value_json=3.0,
        page_number=1,
        confidence=0.98,
    )
    session.add(obs)
    session.flush()

    obs_replacement = Observation(
        certificate_id=cert.id,
        product_id=prod.id,
        field_path="thickness_mm",
        source_text="3.05 mm",
        raw_value_json="3.05 mm",
        normalized_value_json=3.05,
        page_number=1,
        confidence=0.99,
    )
    session.add(obs_replacement)
    session.flush()

    rs = session.scalar(select(RuleSet).order_by(RuleSet.id).limit(1))
    assert rs is not None
    run = ClassificationRun(
        certificate_id=cert.id,
        rule_set_id=rs.id,
        approval_status="approved",
        input_snapshot_json={},
        demo_notice=settings.demo_notice,
    )
    session.add(run)
    session.flush()

    res = ClassificationResult(
        classification_run_id=run.id,
        product_id=prod.id,
        outcome="classified",
        fraction="72091504",
        nico="01",
    )
    session.add(res)
    session.flush()

    c1 = ClassificationCandidate(
        classification_result_id=res.id,
        rank=1,
        fraction="72091504",
        nico="01",
        description="Bobinas laminadas en caliente e >= 3mm",
        support_level="fully_supported",
    )
    c2 = ClassificationCandidate(
        classification_result_id=res.id,
        rank=2,
        fraction="72091504",
        nico="02",
        description="Bobinas laminadas e < 3mm",
        support_level="conditional",
    )
    c3 = ClassificationCandidate(
        classification_result_id=res.id,
        rank=3,
        fraction="72091599",
        nico="00",
        description="Otras bobinas",
        support_level="conditional",
    )
    session.add_all([c1, c2, c3])
    session.flush()

    factor = CandidateFactor(
        candidate_id=c1.id,
        sequence=1,
        rule_code="chapter72.thickness.gte_3mm",
        outcome="matched",
        operator=">=",
        expected_json={"min": 3.0},
        observed_json={"value": 3.0},
        explanation="Espesor 3.0 mm cumple límite inferior",
        required_for_selection=True,
    )
    session.add(factor)
    session.flush()

    ev_link = EvidenceLink(
        candidate_factor_id=factor.id,
        observation_id=obs.id,
        source_type="observation",
        field_path="thickness_mm",
        source_reference_json={"page": 1},
    )
    session.add(ev_link)
    session.flush()

    sel = ClassificationSelection(
        classification_result_id=res.id,
        candidate_id=c1.id,
        person_name="Ing. Auditor",
        reason="Especificaciones geométricas y químicas comprobadas",
        workstation_name="WS-AUDIT",
        created_at=datetime.now(timezone.utc),
    )
    session.add(sel)
    session.flush()

    appr = ApprovalEvent(
        classification_run_id=run.id,
        from_status="needs_review",
        to_status="approved",
        person_name="Jefe Aduanal",
        reason="Aprobación definitiva de clasificación",
        workstation_name="WS-SUPERVISOR",
        created_at=datetime.now(timezone.utc),
    )
    session.add(appr)
    session.flush()

    corr = Correction(
        certificate_id=cert.id,
        previous_observation_id=obs.id,
        replacement_observation_id=obs_replacement.id,
        person_name="Auditor Calidad",
        reason="Ajuste por redondeo del vernier",
        workstation_name="WS-AUDIT",
        created_at=datetime.now(timezone.utc),
    )
    session.add(corr)
    session.flush()

    # --- 1. Export as official=True ---
    dest_path = tmp_path / "report_official.xlsx"
    service = ExcelExportService(settings)
    out_path = service.export_certificates(
        session,
        certificate_ids=[cert.id],
        destination=dest_path,
        official=True,
    )
    assert out_path.is_file()

    wb = openpyxl.load_workbook(out_path)
    expected_sheets = [
        "Resumen", "Actas", "Coladas", "Rollos", "Composición",
        "Clasificación", "Evidencia", "Auditoría",
    ]
    assert wb.sheetnames == expected_sheets

    # Verify Resumen sheet
    ws_resumen = wb["Resumen"]
    a1_val = str(ws_resumen["A2"].value)
    assert settings.demo_notice in a1_val
    assert "PRELIMINAR" not in a1_val

    # Verify Resumen hyperlinks to other sheets
    hyperlinks = [cell.hyperlink.target for row in ws_resumen.iter_rows() for cell in row if cell.hyperlink]
    assert any("Actas" in h for h in hyperlinks)
    assert any("Clasificación" in h for h in hyperlinks)
    assert any("Auditoría" in h for h in hyperlinks)

    # Verify Clasificación sheet
    ws_clas = wb["Clasificación"]
    headers = [cell.value for cell in ws_clas[3]]
    assert "candidato_elegido" in headers
    assert "alternativa_2" in headers
    assert "alternativa_3" in headers
    assert "seleccionado_por" in headers
    assert "factores_clave" in headers
    assert "evidencia_enlace" in headers

    row4 = {ws_clas.cell(3, col).value: ws_clas.cell(4, col).value for col in range(1, len(headers) + 1)}
    assert "72091504-01" in str(row4["candidato_elegido"])
    assert "72091504-02" in str(row4["alternativa_2"])
    assert "72091599-00" in str(row4["alternativa_3"])
    assert row4["seleccionado_por"] == "Ing. Auditor"
    assert row4["fracción"] == "72091504"
    assert row4["NICO"] == "01"
    assert row4["aprobación_fracción_NICO"] == "Aprobado"
    assert ws_clas.sheet_view.showGridLines is False
    assert ws_clas.freeze_panes == "E4"
    assert len(ws_clas.conditional_formatting) > 0
    assert wb["Actas"]["M4"].value == "Aprobado"
    assert wb["Rollos"]["O4"].value is None
    assert "chapter72.thickness.gte_3mm" in str(row4["factores_clave"])

    # Check internal hyperlink in Clasificación pointing to Evidencia
    ev_col_idx = headers.index("evidencia_enlace") + 1
    ev_cell = ws_clas.cell(4, ev_col_idx)
    assert ev_cell.hyperlink is not None
    assert "Evidencia" in ev_cell.hyperlink.target

    # Verify Evidencia sheet
    ws_ev = wb["Evidencia"]
    ev_headers = [cell.value for cell in ws_ev[3]]
    assert "campo" in ev_headers
    assert "regla_factor" in ev_headers
    regla_col_idx = ev_headers.index("regla_factor") + 1
    assert any(
        ws_ev.cell(r, regla_col_idx).value == "chapter72.thickness.gte_3mm"
        for r in range(4, ws_ev.max_row + 1)
    )

    # Verify Auditoría sheet
    ws_aud = wb["Auditoría"]
    aud_headers = [cell.value for cell in ws_aud[3]]
    assert "tipo_evento" in aud_headers
    assert "persona" in aud_headers
    assert "motivo" in aud_headers
    assert "vigencia" in aud_headers

    tipo_col_idx = aud_headers.index("tipo_evento") + 1
    event_types = [ws_aud.cell(r, tipo_col_idx).value for r in range(4, ws_aud.max_row + 1)]
    assert "SELECCION_CANDIDATO" in event_types
    assert "APROBACION_ESTADO" in event_types
    assert "CORRECCION_DATO" in event_types

    # A selected rank-2 candidate must not also appear as an alternative.
    session.add(ClassificationSelection(
        classification_result_id=res.id, candidate_id=c2.id, supersedes_selection_id=sel.id,
        person_name="Supervisor", reason="Revisión de alternativa", workstation_name="WS-TEST"))
    extra_factor = CandidateFactor(candidate_id=c2.id, sequence=1, rule_code="second.rule",
        outcome="matched", explanation="Segunda regla", required_for_selection=True)
    session.add(extra_factor)
    session.flush()
    session.add(EvidenceLink(candidate_factor_id=extra_factor.id, observation_id=obs.id,
        source_type="observation", field_path="thickness_mm", source_reference_json={}))
    session.add(ApprovalEvent(classification_run_id=run.id, from_status="approved", to_status="needs_review",
        person_name="Supervisor", reason="Nueva revisión", workstation_name="WS-TEST",
        created_at=datetime.now(timezone.utc)))
    obs_replacement.is_current = False
    session.flush()

    # --- 2. Export as official=False ---
    dest_prelim = tmp_path / "report_prelim.xlsx"
    out_prelim = service.export_certificates(
        session,
        certificate_ids=[cert.id],
        destination=dest_prelim,
        official=False,
    )
    wb_prelim = openpyxl.load_workbook(out_prelim)
    ws_prelim_resumen = wb_prelim["Resumen"]
    assert "REPORTE DE CONSULTA" in str(ws_prelim_resumen["A2"].value)
    assert wb_prelim["Clasificación"]["R4"].value == "02"
    assert wb_prelim["Clasificación"]["E4"].value == "72091504-02"
    assert wb_prelim["Clasificación"]["F4"].value == "72091504-01"
    assert wb_prelim["Clasificación"]["G4"].value == "72091599-00"
    assert wb_prelim["Evidencia"]["M4"].value == "chapter72.thickness.gte_3mm; second.rule"
    audit = list(wb_prelim["Auditoría"].iter_rows(min_row=4, values_only=True))
    assert [r for r in audit if r[0] == "CORRECCION_DATO"][0][-1] == "Reemplazada"
    assert [r for r in audit if r[0] == "APROBACION_ESTADO" and r[1] == appr.id][0][-1] == "Reemplazada"
    wb.close()
    wb_prelim.close()

    res.details_json = {"selection_cleared": True}
    res.fraction = None
    res.nico = None
    res.outcome = "needs_review"
    run.approval_status = "needs_review"
    session.add(Product(certificate_id=cert.id, product_identifier="SIN-CLASIFICACION"))
    session.flush()
    service.export_certificates(session, certificate_ids=[cert.id], destination=dest_prelim, official=False)
    cleared = openpyxl.load_workbook(dest_prelim)
    assert cleared["Clasificación"]["E4"].value == "Pendiente de selección"
    assert cleared["Clasificación"]["Q4"].value is None
    assert cleared["Clasificación"]["S4"].value == "Pendiente de aprobación"
    assert cleared["Clasificación"]["S5"].value == "Pendiente de clasificación"
    selection_rows = [r for r in cleared["Auditoría"].iter_rows(min_row=4, values_only=True) if r[0] == "SELECCION_CANDIDATO"]
    assert all(r[-1] == "Retirada" for r in selection_rows)
    cleared.close()


def test_official_export_blocks_unapproved_records(db_session, tmp_path):
    session, db_url = db_session
    today = date.today()
    settings = Settings(database_url=db_url, storage_root=tmp_path / "storage")

    _, cert, _ = _seed_test_document_and_cert(session, "CERT-UNAPPROVED", today, status="needs_review")

    dest_path = tmp_path / "report_unapproved.xlsx"
    service = ExcelExportService(settings)

    with pytest.raises(ValueError) as exc_info:
        service.export_certificates(
            session,
            certificate_ids=[cert.id],
            destination=dest_path,
            official=True,
        )
    assert "reporte oficial sólo admite actas aprobadas" in str(exc_info.value)


def test_api_certificates_and_export_endpoints(db_session, tmp_path):
    session, db_url = db_session
    today = date.today()
    settings = Settings(
        database_url=db_url,
        test_database_url=db_url,
        storage_root=tmp_path / "storage",
        backup_root=tmp_path / "backups",
    )

    _, cert_approved, _ = _seed_test_document_and_cert(session, "CERT-API-APP", today, status="approved")
    _, cert_draft, _ = _seed_test_document_and_cert(session, "CERT-API-DRF", today, status="draft")

    app = create_app(settings)

    def _override_get_session():
        yield session

    app.dependency_overrides[get_session] = _override_get_session

    with TestClient(app) as client:
        # 1. Test GET /api/v1/certificates with filters
        res = client.get("/api/v1/certificates?period=today&processing_status=succeeded")
        assert res.status_code == 200
        body = res.json()
        assert "filters" in body["meta"]
        assert body["meta"]["filters"]["period"] == "today"
        assert body["meta"]["filters"]["processing_status"] == "succeeded"

        # 2. Test POST /api/v1/exports with official=True on draft cert -> 409
        res_exp_err = client.post(
            "/api/v1/exports",
            json={
                "person_name": "Auditor",
                "certificate_ids": [cert_draft.id],
                "official": True,
            },
        )
        assert res_exp_err.status_code == 409
        assert res_exp_err.json()["error"]["code"] == "official_export_requires_approval"

        # 3. Test POST /api/v1/exports with official=False -> 202
        res_exp_ok = client.post(
            "/api/v1/exports",
            json={
                "person_name": "Auditor",
                "workstation_name": "WS-TEST",
                "certificate_ids": [cert_draft.id],
                "official": False,
                "filters": {"period": "month"},
            },
        )
        assert res_exp_ok.status_code == 202
        exp_id = res_exp_ok.json()["data"]["export_id"]

        # 4. Test GET /api/v1/exports/{export_id}
        res_status = client.get(f"/api/v1/exports/{exp_id}")
        assert res_status.status_code == 200
        data = res_status.json()["data"]
        assert data["id"] == exp_id
        assert data["person_name"] == "Auditor"
        assert data["workstation_name"] == "WS-TEST"
        assert data["filters"] == {"period": "month"}
        assert data["status"] in {"queued", "running", "succeeded"}
        contract = client.get("/openapi.json").json()
        assert contract["paths"]["/api/v1/exports"]["post"]["responses"]["202"]["content"]["application/json"]["schema"]["$ref"].endswith("ExportCreatedEnvelope")
        assert contract["components"]["schemas"]["ExportStatusRead"]["properties"]["status"]["enum"] == ["queued", "running", "succeeded", "failed"]


@pytest.mark.parametrize("scenario", ["missing", "latest_unapproved", "partial_explicit"])
def test_official_export_requires_approved_run_for_every_certificate(db_session, tmp_path, scenario):
    session, db_url = db_session
    settings = Settings(database_url=db_url, test_database_url=db_url, storage_root=tmp_path / "storage")
    _, cert, _ = _seed_test_document_and_cert(session, "OFFICIAL", date.today(), status="approved")
    _, missing_cert, _ = _seed_test_document_and_cert(session, "MISSING", date.today(), status="approved")
    rs = session.scalar(select(RuleSet).order_by(RuleSet.id).limit(1))
    run = ClassificationRun(certificate_id=cert.id, rule_set_id=rs.id, approval_status="approved",
                            input_snapshot_json={}, demo_notice=settings.demo_notice)
    session.add(run)
    session.flush()
    ids = [cert.id, missing_cert.id] if scenario != "latest_unapproved" else [cert.id]
    explicit = [run.id] if scenario == "partial_explicit" else None
    if scenario == "latest_unapproved":
        session.add(ClassificationRun(certificate_id=cert.id, rule_set_id=rs.id, approval_status="needs_review",
                                     input_snapshot_json={}, demo_notice=settings.demo_notice))
        session.flush()
    with pytest.raises(ValueError):
        ExcelExportService(settings).export_certificates(session, certificate_ids=ids,
            classification_run_ids=explicit, destination=tmp_path / "blocked.xlsx", official=True)
    assert not (tmp_path / "blocked.xlsx").exists()
    app = create_app(settings)
    def override_session():
        yield session
    app.dependency_overrides[get_session] = override_session
    with TestClient(app) as client:
        response = client.post("/api/v1/exports", json={"certificate_ids": ids,
            "classification_run_ids": explicit or [], "official": True, "person_name": "Auditor"})
        assert response.status_code == 409
        assert response.json()["error"]["code"] == "official_export_requires_approval"
        assert "409" in client.get("/openapi.json").json()["paths"]["/api/v1/exports"]["post"]["responses"]


def test_heat_export_limits_results_and_audit_and_never_links_unrelated_evidence(db_session, tmp_path):
    session, db_url = db_session
    settings = Settings(database_url=db_url, storage_root=tmp_path / "storage")
    _, cert, _ = _seed_test_document_and_cert(session, "HEAT-SCOPE", date.today())
    heats = [Heat(certificate_id=cert.id, heat_no=f"HEAT-{i}") for i in range(2)]
    session.add_all(heats)
    session.flush()
    products = [Product(certificate_id=cert.id, heat_id=h.id, product_identifier=f"COIL-{i}")
                for i, h in enumerate(heats)]
    session.add_all(products)
    session.flush()
    rs = session.scalar(select(RuleSet).order_by(RuleSet.id).limit(1))
    run = ClassificationRun(certificate_id=cert.id, rule_set_id=rs.id, approval_status="draft",
                            input_snapshot_json={}, demo_notice=settings.demo_notice)
    session.add(run)
    session.flush()
    session.add_all([ClassificationResult(classification_run_id=run.id, product_id=p.id,
        outcome="classified", fraction="72091504", nico="01") for p in products])
    unrelated = Observation(certificate_id=cert.id, product_id=products[1].id,
        field_path="thickness_mm", raw_value_json=3, normalized_value_json=3)
    session.add(unrelated)
    session.flush()
    session.add(Correction(certificate_id=cert.id, replacement_observation_id=unrelated.id,
        person_name="Auditor", reason="Corrección ajena", workstation_name="WS-TEST"))
    session.flush()
    path = ExcelExportService(settings).export_certificates(session, certificate_ids=[cert.id],
        heat_ids=[heats[0].id], destination=tmp_path / "heat.xlsx", official=False)
    workbook = openpyxl.load_workbook(path)
    assert workbook["Clasificación"].max_row == 4
    assert workbook["Clasificación"]["C4"].value == products[0].id
    assert workbook["Clasificación"]["L4"].value == "Sin evidencia"
    assert workbook["Clasificación"]["L4"].hyperlink is None
    assert workbook["Auditoría"].max_row == 3
    workbook.close()
    general = Observation(certificate_id=cert.id, field_path="certificate_no", raw_value_json="HEAT-SCOPE")
    inherited = Observation(certificate_id=cert.id, heat_id=heats[0].id,
                            field_path="composition.C", normalized_value_json=0.15)
    session.add_all([general, inherited])
    session.flush()
    ExcelExportService(settings).export_certificates(session, certificate_ids=[cert.id],
        heat_ids=[heats[0].id], destination=path, official=False)
    workbook = openpyxl.load_workbook(path)
    assert workbook["Evidencia"].max_row == 5
    assert workbook["Clasificación"]["L4"].hyperlink.target == "#'Evidencia'!A5"
    assert workbook["Clasificación"]["I4"].value == "Resolución unívoca pendiente de aprobación"
    workbook.close()


def test_export_request_pins_latest_run(db_session, tmp_path):
    session, db_url = db_session
    settings = Settings(database_url=db_url, test_database_url=db_url, storage_root=tmp_path / "storage")
    _, cert, _ = _seed_test_document_and_cert(session, "PINNED", date.today(), status="approved")
    rs = session.scalar(select(RuleSet).order_by(RuleSet.id).limit(1))
    run = ClassificationRun(certificate_id=cert.id, rule_set_id=rs.id, approval_status="approved",
                            input_snapshot_json={}, demo_notice=settings.demo_notice)
    session.add(run)
    session.flush()
    app = create_app(settings)
    def override_session():
        yield session
    app.dependency_overrides[get_session] = override_session
    with TestClient(app) as client:
        response = client.post("/api/v1/exports", json={"certificate_ids": [cert.id],
            "official": True, "person_name": "Auditor"})
        assert response.status_code == 202
        export = session.get(Export, response.json()["data"]["export_id"])
        assert export.scope_json["classification_run_ids"] == [run.id]
        assert client.get("/api/v1/certificates?processing_status=unknown").status_code == 422
        invalid_range = client.get("/api/v1/certificates?date_from=2026-10-08&date_to=2026-10-07")
        assert invalid_range.status_code == 400
        assert invalid_range.json()["error"]["code"] == "invalid_date_range"
    session.add(ClassificationRun(certificate_id=cert.id, rule_set_id=rs.id, approval_status="draft",
                                 input_snapshot_json={}, demo_notice=settings.demo_notice))
    session.flush()
    pinned = ExcelExportService.resolve_runs(session, [cert.id], export.scope_json["classification_run_ids"], True)
    assert [item.id for item in pinned] == [run.id]
    assert ExcelExportService.resolve_runs(session, [cert.id], [], False) == []
